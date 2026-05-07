from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Create a new Document
doc = Document()

# Add Title
title = doc.add_heading('Design Rationale: Campus Concierge Web Portal', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Add Introduction
doc.add_heading('1. Executive Summary', level=1)
doc.add_paragraph(
    "The Campus Concierge portal is designed as a unified, modern, and highly interactive digital hub for university students. "
    "The primary goal of the design is to consolidate fragmented university services (LMS, Qalam, S3C services) into a single, cohesive user interface. "
    "By adopting a mobile-first responsive architecture and an engaging visual language, the system reduces cognitive load, improves accessibility, "
    "and creates a premium user experience tailored specifically for the NUST student body."
)

# Add Aesthetic Choices
doc.add_heading('2. Visual Identity & Aesthetics', level=1)
p1 = doc.add_paragraph()
p1.add_run('Color Palette: ').bold = True
p1.add_run(
    "The application heavily utilizes NUST's core colors (deep blues like #0B2F6B and accent yellows like #FACC15). "
    "This creates an immediate sense of institutional belonging. Soft background gradients and 'glassmorphism' (blurred translucent backgrounds) "
    "add depth and a modern 'app-like' feel to the web experience."
)

p2 = doc.add_paragraph()
p2.add_run('Typography: ').bold = True
p2.add_run(
    "The UI employs a dual-font strategy using 'Public Sans' and 'Space Grotesk'. This combination ensures maximum legibility for functional elements "
    "while allowing headers and branding to stand out with a distinct, contemporary character."
)

# Add Layout and System
doc.add_heading('3. Bento Card System & Surface Language', level=1)
doc.add_paragraph(
    "The interface is anchored by a 'Bento Card System', which organizes information into modular, rounded (22px radius) containers. "
    "This modularity ensures that the layout is highly adaptable across different screen sizes. Interactive elements like cards have subtle "
    "hover states (e.g., transforming along the Y-axis and increasing shadow depth) to provide immediate tactile feedback, making the interface feel alive."
)

# Add UX & Interactions
doc.add_heading('4. Micro-interactions and Feedback', level=1)
doc.add_paragraph(
    "Using Framer Motion, the application implements fluid micro-interactions. Examples include the dynamic 'active tab' indicator in the navigation bar, "
    "spring-based physics for the mobile slide-over menus, and gesture-driven bottom sheets for notifications. These animations are not merely decorative; "
    "they guide the user's focus and provide context about spatial relationships within the app."
)

# Add Responsive Design
doc.add_heading('5. Mobile-First Responsiveness', level=1)
doc.add_paragraph(
    "Recognizing that most students access the portal via smartphones, the design inherently prioritizes mobile usability. "
    "Complex desktop components gracefully degrade into touch-friendly alternatives. For instance, the desktop notification dropdown becomes an intuitive, "
    "swipeable bottom sheet on mobile, and the top navigation collapses into an easily accessible hamburger menu."
)

# Add Theming
doc.add_heading('6. Adaptive Theming', level=1)
doc.add_paragraph(
    "The UI supports comprehensive theming (Light, Dark, and a custom 'NUST Exclusive' theme) managed via CSS custom properties. "
    "This allows the app to adapt to user preferences and ambient lighting conditions, reducing eye strain and improving overall comfort during extended use."
)

# Save the document
doc.save('Campus_Concierge_Design_Rationale.docx')
print("Document generated successfully.")
